# -*- coding: utf-8 -*-
"""Gộp các file 'Auto Task.pending-*.xlsx' vào 'Auto Task.xlsx'.
- Mỗi pending chứa đúng 1 sheet (header + các dòng dữ liệu) -> append vào sheet cùng tên ở file gốc.
- Bỏ qua dòng trùng (so theo time_start+time_end+title).
- Gộp xong: đổi tên pending thành '.merged' (không xoá hẳn cho an toàn).
- Nếu file gốc đang khoá (mở trong Excel) -> báo rõ, KHÔNG đổi tên pending.
"""
import glob
import os
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from paths import auto_task_xlsx, desktop, station  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DESK = desktop()
XLSX = auto_task_xlsx()
# Bản pending có thể nằm ở 4 nơi: trong trạm (mặc định TỪ 21/09/2026), cạnh script
# (mặc định cũ, khi engine còn nằm ngoài trạm), cạnh sổ, hoặc trên Desktop (các bản
# cũ, hoặc khi Excel đang khoá file). GIỮ CẢ BỐN: mỗi chỗ tương ứng một thời kỳ, bỏ
# ra là mất luôn những dòng chưa gộp sinh ra trong thời kỳ đó.
PENDING_DIRS = [os.path.join(station(), "_xlsx-pending"),
                os.path.join(HERE, "_xlsx-pending"),
                os.path.dirname(XLSX), DESK]
COLS = ["time_start", "time_end", "status", "script", "title", "video_location", "post", "note"]


def ensure_sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(title=name)
    ws.append(COLS)
    for c in range(1, len(COLS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for i, w in enumerate([19, 19, 10, 34, 52, 16, 60, 50], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    return ws


def existing_keys(ws):
    keys = set()
    for r in range(2, ws.max_row + 1):
        keys.add((ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 5).value))
    return keys


def copy_row(src_ws, src_r, dst_ws):
    r = dst_ws.max_row + 1
    for c in range(1, len(COLS) + 1):
        s = src_ws.cell(src_r, c)
        d = dst_ws.cell(r, c)
        d.value = s.value
        if s.hyperlink:
            d.hyperlink = s.hyperlink.target if hasattr(s.hyperlink, "target") else s.hyperlink
            d.style = "Hyperlink"


def main():
    pend = []
    for d in PENDING_DIRS:
        pend.extend(glob.glob(os.path.join(d, "Auto Task.pending-*.xlsx")))
    pend = sorted(set(pend))
    if not pend:
        print(f"Không có file pending nào (đã tìm ở: {'; '.join(PENDING_DIRS)}).")
        return 0
    print(f"Tìm thấy {len(pend)} file pending.")

    if not os.path.isfile(XLSX):
        print(f"LỖI: không thấy file gốc {XLSX}")
        return 1
    try:
        wb = load_workbook(XLSX)
    except PermissionError:
        print("LOCKED: file gốc Auto Task.xlsx đang MỞ trong Excel -> đóng Excel rồi chạy lại.")
        return 2

    added = 0
    merged_files = []
    for pf in pend:
        try:
            pwb = load_workbook(pf)
        except Exception as e:
            print(f"  bỏ qua {os.path.basename(pf)}: {type(e).__name__}")
            continue
        any_row = False
        for sname in pwb.sheetnames:
            pws = pwb[sname]
            dst = ensure_sheet(wb, sname)
            keys = existing_keys(dst)
            for r in range(2, pws.max_row + 1):
                k = (pws.cell(r, 1).value, pws.cell(r, 2).value, pws.cell(r, 5).value)
                if all(v is None for v in k):
                    continue
                if k in keys:
                    print(f"  trùng (bỏ): {sname} | {k[2]}")
                    continue
                copy_row(pws, r, dst)
                keys.add(k)
                added += 1
                any_row = True
                print(f"  + {sname} | {k[2]}")
        if any_row:
            merged_files.append(pf)

    if added == 0:
        print("Không có dòng mới để gộp (tất cả đã trùng).")
    try:
        wb.save(XLSX)
    except PermissionError:
        print("LOCKED khi lưu: file gốc đang mở trong Excel -> đóng Excel rồi chạy lại. Chưa đổi tên pending.")
        return 2
    print(f"ĐÃ GỘP {added} dòng vào {XLSX}")

    for pf in merged_files:
        try:
            os.rename(pf, pf + ".merged")
            print(f"  đổi tên -> {os.path.basename(pf)}.merged")
        except OSError as e:
            print(f"  không đổi tên được {os.path.basename(pf)}: {e}")
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    raise SystemExit(main())
